"""
Book Scraper - Portfolio Project
Author: Oscar J. Villa García
Description: Scrapes books.toscrape.com and exports data to Excel with formatting.
"""

import requests
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
import time
import logging
from datetime import datetime
from pathlib import Path

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger(__name__)

BASE_URL   = "https://books.toscrape.com/catalogue/"
START_URL  = "https://books.toscrape.com/catalogue/page-1.html"
MAX_PAGES  = 10
DELAY      = 0.5
OUTPUT_DIR = Path(__file__).parent.parent / "output"
RATING_MAP = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

def get_soup(url):
    headers = {"User-Agent": "Mozilla/5.0 (portfolio-scraper/1.0)"}
    response = requests.get(url, headers=headers, timeout=10)
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser")

def parse_books(soup):
    books = []
    for article in soup.select("article.product_pod"):
        title  = article.h3.a["title"]
        price  = float(article.select_one(".price_color").text.strip("£Â"))
        rating = RATING_MAP.get(article.p["class"][1], 0)
        avail  = article.select_one(".availability").text.strip()
        link   = BASE_URL + article.h3.a["href"].replace("../", "")
        books.append({"Title": title, "Price (£)": price, "Rating": rating, "Availability": avail, "URL": link})
    return books

def scrape(max_pages=MAX_PAGES):
    all_books, url, page = [], START_URL, 1
    while url:
        log.info(f"Scraping page {page}")
        soup = get_soup(url)
        all_books.extend(parse_books(soup))
        if max_pages and page >= max_pages:
            break
        next_btn = soup.select_one("li.next a")
        url = BASE_URL + next_btn["href"] if next_btn else None
        page += 1
        time.sleep(DELAY)
    log.info(f"Done. {len(all_books)} books scraped.")
    return all_books

def export_to_excel(books):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(books)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = OUTPUT_DIR / f"books_report_{ts}.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Books")
        ws = writer.book.active
        fill = PatternFill("solid", fgColor="1F4E79")
        font = Font(bold=True, color="FFFFFF", size=11)
        for col in range(1, len(df.columns)+1):
            cell = ws.cell(row=1, column=col)
            cell.fill = fill
            cell.font = font
        ws.freeze_panes = "A2"
        for col in range(1, len(df.columns)+1):
            ws.column_dimensions[get_column_letter(col)].width = 20
    log.info(f"Excel saved -> {out_path}")
    return out_path

if __name__ == "__main__":
    books = scrape()
    out_path = export_to_excel(books)
    print(f"\n✅ Report ready: {out_path}")
